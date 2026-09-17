from pathlib import Path

from docx import Document
from pptx import Presentation


INPUT_DIR = Path("input")


def create_docx(filename: str, title: str, body: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(body)
    document.save(INPUT_DIR / filename)


def create_pptx(filename: str, title: str, body: str) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    presentation.save(INPUT_DIR / filename)


def main() -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)

    create_pptx(
        "general_request.pptx",
        "Architecture Summary",
        "Analyze this architecture and summarize the main components and responsibilities.",
    )

    create_docx(
        "coding_request.docx",
        "Python Automation",
        "Create a Python automation script to validate employee records and report invalid records.",
    )

    create_docx(
        "sql_request.docx",
        "Employee Query",
        "Create a SQL query to find employees whose salary is greater than 100000.",
    )

    create_docx(
        "unsafe_sql_request.docx",
        "Inactive Customers",
        "Delete all inactive customer records from the customers table.",
    )

    print("Demo input files created in the input folder.")


if __name__ == "__main__":
    main()
