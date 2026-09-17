from pathlib import Path

from docx import Document
from pptx import Presentation


class DocumentProcessingError(Exception):
    """Raised when a supported document cannot be processed."""


def extract_text(file_path: str, max_chars: int = 50000) -> str:
    """Extract readable text from a DOCX or PPTX file, capped at max_chars."""
    path = Path(file_path)

    if not path.exists():
        raise DocumentProcessingError(f"Input file not found: {path}")

    suffix = path.suffix.lower()

    try:
        if suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".pptx":
            text = _extract_pptx(path)
        else:
            raise DocumentProcessingError(
                f"Unsupported file type '{suffix}'. Use .docx or .pptx."
            )

        if not text.strip():
            raise DocumentProcessingError(f"No readable text found in '{path.name}'.")

        return text[:max_chars]
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError(
            f"Failed to extract text from '{path.name}': {exc}"
        ) from exc


def _extract_docx(path: Path) -> str:
    document = Document(path)
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            sections.append(" | ".join(values))

    return "\n".join(sections).strip()


def _extract_pptx(path: Path) -> str:
    presentation = Presentation(path)
    sections = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

        if slide_text:
            sections.append(f"[Slide {slide_number}]\n" + "\n".join(slide_text))

    return "\n".join(sections).strip()
