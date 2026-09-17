import csv
import json
from pathlib import Path

from openpyxl import Workbook


class ReportingError(Exception):
    """Raised when report generation fails."""


CALL_LEDGER_HEADERS = [
    "run_id",
    "input_file",
    "phase",
    "provider",
    "model",
    "started_at",
    "duration_seconds",
    "outcome",
]


def write_audit_log(records: list[dict], output_file: str) -> None:
    """Write one JSON object per line so the audit file is easy to inspect."""
    try:
        path = Path(output_file)
        path.parent.mkdir(parents=True, exist_ok=True)

        with path.open("w", encoding="utf-8") as file:
            for record in records:
                file.write(json.dumps(record, ensure_ascii=False) + "\n")
    except OSError as exc:
        raise ReportingError(f"Unable to write audit log: {exc}") from exc


def append_call_ledger(calls: list[dict], output_file: str) -> None:
    """Append every individual LLM call's timing to a persistent CSV ledger.

    Unlike the audit log and Excel report (which are rewritten each run), this
    file is appended to across runs so call durations can be tracked/recorded
    over time — one row per real LLM call (generation, guardrail, evaluation,
    including a failed Gemini attempt before a fallback).
    """
    try:
        path = Path(output_file)
        path.parent.mkdir(parents=True, exist_ok=True)
        write_header = not path.exists()

        with path.open("a", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=CALL_LEDGER_HEADERS)
            if write_header:
                writer.writeheader()
            for call in calls:
                writer.writerow({header: call.get(header, "") for header in CALL_LEDGER_HEADERS})
    except OSError as exc:
        raise ReportingError(f"Unable to write LLM call ledger: {exc}") from exc


def write_excel_report(
    records: list[dict],
    output_file: str,
    run_metadata: dict | None = None,
    calls: list[dict] | None = None,
) -> None:
    """Create a governance report with a run summary, configuration, audit trail and call ledger."""
    try:
        path = Path(output_file)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"

        total = len(records)
        decision_counts = {
            decision: sum(record.get("final_decision") == decision for record in records)
            for decision in ("PASS", "BLOCK", "MANUAL_REVIEW", "TECHNICAL_ERROR")
        }

        summary.append(["Metric", "Value"])
        summary.append(["Run ID", (run_metadata or {}).get("run_id", "")])
        summary.append(["Total Requests", total])
        for decision, count in decision_counts.items():
            summary.append([decision.title().replace("_", " "), count])

        if run_metadata:
            configuration = workbook.create_sheet("Configuration")
            configuration.append(["Key", "Value"])
            for key, value in run_metadata.items():
                configuration.append([key, value])

        details = workbook.create_sheet("Audit")
        headers = [
            "run_id",
            "timestamp",
            "input_file",
            "intent",
            "route_reason",
            "model",
            "generation_provider",
            "generation_latency_seconds",
            "guardrail_decision",
            "guardrail_risk",
            "guardrail_source",
            "guardrail_latency_seconds",
            "evaluation_verdict",
            "evaluation_score",
            "evaluation_latency_seconds",
            "total_latency_seconds",
            "final_decision",
            "status",
            "failure_type",
            "reason",
        ]
        details.append(headers)

        for record in records:
            details.append([record.get(header, "") for header in headers])

        if calls:
            call_sheet = workbook.create_sheet("LLM Calls")
            call_sheet.append(CALL_LEDGER_HEADERS)
            for call in calls:
                call_sheet.append([call.get(header, "") for header in CALL_LEDGER_HEADERS])

        workbook.save(path)
    except OSError as exc:
        raise ReportingError(f"Unable to write Excel report: {exc}") from exc
