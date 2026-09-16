from openpyxl import load_workbook

from reporting.excel import save_report


def test_failures_sheet_uses_each_rows_own_metric_threshold(tmp_path):
    # detail_rows ends on a "bias" row so that, under the pre-fix bug, the
    # Failures sheet would reuse "bias" thresholds (pass=0.80, review=0.60)
    # for every row instead of each row's own metric.
    detail_rows = [
        {
            "generator": "gen1", "judge": "judge1", "test_id": "t1",
            "metric": "bias", "score": 0.95, "verdict": "PASS",
            "passed": True, "status": "COMPLETED", "reason": "",
            "error": "", "duration_seconds": 1.0, "llm_call_timings": [],
        },
    ]

    # Same score, different metrics -> different thresholds ->
    # different expected verdicts:
    #   completeness: pass=0.70 -> 0.72 is PASS
    #   hallucination: pass=0.80, review=0.60 -> 0.72 is REVIEW
    failures = [
        {
            "generator": "gen1", "judge": "judge1", "test_id": "t1",
            "metric": "completeness", "score": 0.72, "verdict": "REVIEW",
            "status": "COMPLETED", "reason": "reason-completeness",
            "error": "", "failure_type": "QUALITY ISSUE",
        },
        {
            "generator": "gen1", "judge": "judge1", "test_id": "t1",
            "metric": "hallucination", "score": 0.72, "verdict": "REVIEW",
            "status": "COMPLETED", "reason": "reason-hallucination",
            "error": "", "failure_type": "QUALITY ISSUE",
        },
    ]

    file_path = save_report(
        summary=[],
        testcase_rows=[],
        detail_rows=detail_rows,
        failures=failures,
        configuration=[],
        report_root=tmp_path,
        run_summary=[{"section": "Run ID", "value": "test-run"}],
        run_id="test-run",
    )

    workbook = load_workbook(file_path)
    sheet = workbook["Failures"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    # headers: Generator, Judge, Test Case, Metric, Score (%), Verdict,
    # Failure Type, Status, Reason, Error - Reason is index 8.
    by_reason = {row[8]: row for row in rows}

    assert by_reason["reason-completeness"][5] == "PASS"
    assert by_reason["reason-hallucination"][5] == "REVIEW"


def test_run_summary_and_llm_call_profile_are_written(tmp_path):
    detail_rows = [
        {
            "generator": "gen1",
            "judge": "judge1",
            "test_id": "t1",
            "metric": "faithfulness",
            "score": 0.9,
            "verdict": "PASS",
            "passed": True,
            "status": "COMPLETED",
            "reason": "ok",
            "error": "",
            "duration_seconds": 12.0,
            "llm_call_timings": [
                {"step": "Truths", "duration_seconds": 3.0},
                {"step": "Claims", "duration_seconds": 4.0},
            ],
        }
    ]

    file_path = save_report(
        summary=[],
        testcase_rows=[],
        detail_rows=detail_rows,
        failures=[],
        configuration=[],
        report_root=tmp_path,
        run_summary=[
            {"section": "Run ID", "value": "run-123"},
            {"section": "Total", "value": "12.0s"},
        ],
        run_id="run-123",
    )

    workbook = load_workbook(file_path)

    assert workbook["Run Summary"]["B2"].value == "run-123"
    profile = workbook["LLM Call Profile"]
    rows = list(profile.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert rows[0][4] == 1
    assert rows[0][5] == "Truths"


def test_generator_comparison_blanks_metric_with_incomplete_status(tmp_path):
    # gen1's hallucination has one COMPLETED row and one ERROR row for the
    # same metric - the per-generator average must not be silently
    # computed from just the completed row.
    detail_rows = [
        {
            "generator": "gen1", "judge": "judge1", "test_id": "t1",
            "metric": "hallucination", "score": 0.9, "verdict": "PASS",
            "passed": True, "status": "COMPLETED", "reason": "",
            "error": "", "duration_seconds": 1.0, "llm_call_timings": [],
        },
        {
            "generator": "gen1", "judge": "judge1", "test_id": "t2",
            "metric": "hallucination", "score": None, "verdict": "ERROR",
            "passed": False, "status": "ERROR", "reason": "",
            "error": "timeout", "duration_seconds": "", "llm_call_timings": [],
        },
    ]

    file_path = save_report(
        summary=[{
            "generator": "gen1", "overall_score": None,
            "passed_testcases": 1, "total_testcases": 2,
            "quality_gate": "FAIL", "technical_error_count": 1,
            "quality_fail_count": 0, "quality_review_count": 0,
        }],
        testcase_rows=[],
        detail_rows=detail_rows,
        failures=[],
        configuration=[],
        report_root=tmp_path,
        run_summary=[{"section": "Run ID", "value": "run-456"}],
        run_id="run-456",
    )

    workbook = load_workbook(file_path)
    sheet = workbook["Generator Comparison"]
    header = [cell.value for cell in sheet[1]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    row = dict(zip(header, rows[0]))

    # openpyxl round-trips a written "" cell value as None after
    # save+reload (XLSX has no distinct empty-string cell representation),
    # so treat either as "blank" rather than asserting the exact type.
    assert not row["Hallucination Score (%)"]
    assert row["Quality Gate"] == "TECHNICAL ERROR"
