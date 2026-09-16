from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config


METRIC_DISPLAY_NAMES = {
    "hallucination": "Hallucination",
    "faithfulness": "Faithfulness",
    "correctness": "Correctness",
    "completeness": "Completeness",
    "answer_relevancy": "Answer Relevancy",
    "bias": "Bias",
}


def _percent(value):
    if value is None or value == "":
        return ""
    return round(float(value) * 100, 1)


def _metric_verdict(score, passed, metric=None):
    if score is None:
        return "ERROR"

    thresholds = config.METRIC_THRESHOLDS.get(
        metric, {"pass": 0.80, "review": 0.60}
    )
    if float(score) >= thresholds["pass"]:
        return "PASS"
    if float(score) >= thresholds["review"]:
        return "REVIEW"
    return "FAIL"


def _metric_interpretation(metric, score, passed):
    if score is None:
        return "Metric evaluation did not complete."

    percent = _percent(score)
    thresholds = config.METRIC_THRESHOLDS.get(
        metric, {"pass": 0.80, "review": 0.60}
    )
    pass_threshold = thresholds["pass"]
    review_threshold = thresholds["review"]

    if passed:
        return (
            f"{percent}% evaluation score - passed the "
            f"{pass_threshold:.2f} quality threshold."
        )

    if float(score) >= review_threshold:
        return (
            f"{percent}% evaluation score - in the REVIEW band "
            f"({review_threshold:.2f} to <{pass_threshold:.2f})."
        )

    return (
        f"{percent}% evaluation score - below the REVIEW band "
        "and marked FAIL."
    )


def _format_sheet(ws):
    ws.freeze_panes = "A2"

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in ws.columns:
        width = 12
        for cell in column_cells:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, 45))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def _write_dataframe(ws, frame):
    if frame.empty:
        ws.append(["No data"])
        return

    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))


def save_generator_responses(response_rows, output_root, run_id=None):
    now = datetime.now()
    output_dir = Path(output_root) / now.strftime("%Y-%m-%d")
    output_dir.mkdir(parents=True, exist_ok=True)

    run_id = run_id or now.strftime("%Y%m%d_%H%M%S")
    file_path = output_dir / f"generator_responses_{run_id}.xlsx"

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        for model, rows in response_rows.items():
            frame = pd.DataFrame(rows)
            sheet_name = str(model).replace(":", "_")[:31]
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

    return file_path


def save_report(
    summary,
    testcase_rows,
    detail_rows,
    failures,
    configuration,
    report_root,
    run_summary=None,
    run_id=None,
):
    now = datetime.now()
    output_dir = Path(report_root) / now.strftime("%Y-%m-%d")
    output_dir.mkdir(parents=True, exist_ok=True)

    run_id = run_id or now.strftime("%Y%m%d_%H%M%S")
    file_path = output_dir / f"generator_comparison_{run_id}.xlsx"

    workbook = Workbook()
    workbook.remove(workbook.active)

    # ---------------------------------------------------------
    # Run Summary - dashboard-style overview, first sheet a viewer sees
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Run Summary")
    ws.append(["Section", "Value"])
    for item in run_summary or []:
        ws.append([item.get("section", ""), item.get("value", "")])
    _format_sheet(ws)

    # ---------------------------------------------------------
    # Executive Summary
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Executive Summary")

    ws.append(["Section", "Value"])
    for item in summary:
        ws.append([item.get("section", ""), item.get("value", "")])

    ws.append([])
    ws.append(
        [
            "Score interpretation",
            "DeepEval scores are displayed as percentages for readability. "
            "A score is an evaluation score, not a literal probability "
            "that a percentage of the response is hallucinated.",
        ]
    )
    ws.append(
        [
            "Hallucination interpretation",
            "For Hallucination, a higher score represents a better "
            "source-grounded / non-hallucinating result. A 0.80 score is "
            "shown as 80% evaluation score and 20 percentage points below "
            "the ideal 100% score; it is not reported as '20% of the "
            "response is hallucinated'.",
        ]
    )

    _format_sheet(ws)

    # ---------------------------------------------------------
    # Generator Comparison
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Generator Comparison")

    detail_frame = pd.DataFrame(detail_rows)

    generators = []
    if not detail_frame.empty and "generator" in detail_frame:
        generators = list(dict.fromkeys(detail_frame["generator"].tolist()))

    metrics = list(METRIC_DISPLAY_NAMES.keys())

    headers = ["Generator"]
    for metric in metrics:
        headers.extend(
            [
                f"{METRIC_DISPLAY_NAMES[metric]} Score (%)",
                f"{METRIC_DISPLAY_NAMES[metric]} Verdict",
            ]
        )
    headers.extend(
        [
            "Overall Score (%)",
            "Testcases Passed",
            "Testcases Failed",
            "Technical Errors",
            "Quality-Fail Testcases",
            "Quality-Review Testcases",
            "Quality Gate",
        ]
    )
    ws.append(headers)

    summary_by_generator = {
        row.get("generator"): row
        for row in summary
        if row.get("generator")
    }

    for generator in generators:
        row = [generator]
        generator_summary = summary_by_generator.get(generator, {})

        for metric in metrics:
            metric_rows = detail_frame[
                detail_frame["generator"].astype(str) == str(generator)
            ]
            metric_rows = metric_rows[metric_rows["metric"] == metric]

            # A metric's per-generator average is blanked out, not
            # silently computed from the successful rows only, if any of
            # this generator's evaluations of that metric didn't complete -
            # the same "don't average over partial failures" rule as
            # weighted_score(), applied at this aggregate too.
            if metric_rows.empty or (
                metric_rows["status"].astype(str) != "COMPLETED"
            ).any():
                score = None
            else:
                scores = pd.to_numeric(
                    metric_rows["score"], errors="coerce"
                ).dropna()
                score = float(scores.mean()) if not scores.empty else None

            passed_values = metric_rows["passed"].tolist()
            passed = bool(passed_values) and all(
                bool(value) for value in passed_values
            )

            row.extend(
                [
                    _percent(score),
                    _metric_verdict(score, passed, metric),
                ]
            )

        overall = generator_summary.get("overall_score")
        passed_count = generator_summary.get("passed_testcases", "")
        total_count = generator_summary.get("total_testcases", "")
        failed_count = (
            total_count - passed_count
            if isinstance(total_count, (int, float))
            and isinstance(passed_count, (int, float))
            else ""
        )
        technical_error_count = generator_summary.get(
            "technical_error_count", 0
        )

        row.extend(
            [
                _percent(overall),
                passed_count,
                failed_count,
                technical_error_count,
                generator_summary.get("quality_fail_count", 0),
                generator_summary.get("quality_review_count", 0),
                (
                    "TECHNICAL ERROR"
                    if technical_error_count
                    else generator_summary.get("quality_gate", "")
                ),
            ]
        )

        ws.append(row)

    _format_sheet(ws)

    # ---------------------------------------------------------
    # Testcase Comparison
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Testcase Comparison")

    if testcase_rows:
        testcase_frame = pd.DataFrame(testcase_rows)

        ordered_columns = ["generator", "test_id"]
        for metric in metrics:
            score_column = f"{metric}_score"
            if score_column in testcase_frame.columns:
                ordered_columns.append(score_column)

        for status_column in ("testcase_status", "testcase_verdict"):
            if status_column in testcase_frame.columns:
                ordered_columns.append(status_column)

        testcase_frame = testcase_frame[
            [column for column in ordered_columns if column in testcase_frame]
        ].copy()

        rename_map = {
            "generator": "Generator",
            "test_id": "Test Case",
            "testcase_status": "Testcase Status",
            "testcase_verdict": "Quality Gate Verdict",
        }

        for metric in metrics:
            rename_map[f"{metric}_score"] = (
                f"{METRIC_DISPLAY_NAMES[metric]} Score (%)"
            )

        testcase_frame.rename(columns=rename_map, inplace=True)

        for column in testcase_frame.columns:
            if column.endswith("Score (%)"):
                testcase_frame[column] = testcase_frame[column].apply(_percent)

        _write_dataframe(ws, testcase_frame)
    else:
        ws.append(["No testcase results"])

    _format_sheet(ws)

    # ---------------------------------------------------------
    # Detailed Metric Reasons
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Detailed Metric Reasons")

    detail_headers = [
        "Generator",
        "Judge",
        "Test Case",
        "Metric",
        "Score (%)",
        "Metric Verdict",
        "Passed Quality Threshold",
        "Gap to Ideal (%)",
        "Reason",
        "Status",
        "Error",
        "Interpretation",
        "Duration (s)",
    ]
    ws.append(detail_headers)

    for item in detail_rows:
        metric = str(item.get("metric", "")).lower()
        score = item.get("score")
        passed = bool(item.get("passed", False))

        if score is not None:
            gap_to_ideal = round((1 - float(score)) * 100, 1)
        else:
            gap_to_ideal = ""

        ws.append(
            [
                item.get("generator", ""),
                item.get("judge", ""),
                item.get("test_id", ""),
                METRIC_DISPLAY_NAMES.get(metric, item.get("metric", "")),
                _percent(score),
                _metric_verdict(score, passed, metric),
                "PASS" if passed else "FAIL",
                gap_to_ideal,
                item.get("reason", ""),
                item.get("status", ""),
                item.get("error", ""),
                _metric_interpretation(metric, score, passed),
                item.get("duration_seconds", ""),
            ]
        )

    _format_sheet(ws)

    # ---------------------------------------------------------
    # Failures
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Failures")

    failure_headers = [
        "Generator",
        "Judge",
        "Test Case",
        "Metric",
        "Score (%)",
        "Verdict",
        "Failure Type",
        "Status",
        "Reason",
        "Error",
    ]
    ws.append(failure_headers)

    for item in failures:
        metric = str(item.get("metric", "")).lower()
        score = item.get("score")

        ws.append(
            [
                item.get("generator", ""),
                item.get("judge", ""),
                item.get("test_id", ""),
                METRIC_DISPLAY_NAMES.get(metric, item.get("metric", "")),
                _percent(score),
                _metric_verdict(score, False, metric),
                item.get("failure_type", ""),
                item.get("status", ""),
                item.get("reason", ""),
                item.get("error", ""),
            ]
        )

    _format_sheet(ws)

    # ---------------------------------------------------------
    # LLM Call Profile - individual internal DeepEval judge-call timing
    # ---------------------------------------------------------
    ws = workbook.create_sheet("LLM Call Profile")
    ws.append(
        [
            "Generator",
            "Judge",
            "Test Case",
            "Metric",
            "Call #",
            "DeepEval Step",
            "Duration (s)",
        ]
    )

    for item in detail_rows:
        for index, call in enumerate(
            item.get("llm_call_timings") or [], start=1
        ):
            ws.append(
                [
                    item.get("generator", ""),
                    item.get("judge", ""),
                    item.get("test_id", ""),
                    METRIC_DISPLAY_NAMES.get(
                        str(item.get("metric", "")).lower(),
                        item.get("metric", ""),
                    ),
                    index,
                    call.get("step", ""),
                    call.get("duration_seconds", ""),
                ]
            )

    _format_sheet(ws)

    # ---------------------------------------------------------
    # Configuration
    # ---------------------------------------------------------
    ws = workbook.create_sheet("Configuration")
    ws.append(["Configuration", "Value"])

    for item in configuration:
        ws.append([item.get("key", ""), item.get("value", "")])

    ws.append([])
    ws.append(
        [
            "Reporting rule",
            "Individual metric scores are displayed as percentages. "
            "Each metric uses its configured PASS and REVIEW thresholds.",
        ]
    )
    ws.append(
        [
            "Testcase rule",
            f"Overall weighted score >= {config.QUALITY_THRESHOLD:.2f} is "
            f"PASS; {config.WARNING_THRESHOLD:.2f} to <{config.QUALITY_THRESHOLD:.2f} "
            f"is REVIEW; below {config.WARNING_THRESHOLD:.2f} is FAIL. "
            "Hallucination, Faithfulness and Correctness must not fall "
            "below their REVIEW thresholds. Technical evaluation errors "
            "are reported separately and never scored as a quality result.",
        ]
    )

    _format_sheet(ws)

    workbook.save(file_path)

    return file_path
